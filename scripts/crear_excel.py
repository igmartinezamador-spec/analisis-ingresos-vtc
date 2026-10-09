"""
Crea analisis_vtc.xlsx a partir de datos/limpios/.

Hojas:
  LEEME          · qué hay en el libro y cómo usarlo
  Resumen        · indicadores y tablas con FÓRMULAS (SUMIFS, COUNTIFS...) + gráficos
  Escenarios     · cálculo del neto: rellenas tus costes (celdas amarillas)
  TD_DiaSemana, TD_Franjas, TD_Producto, TD_Mes_Pago · TABLAS DINÁMICAS reales
  Viajes, Jornadas · los datos (como Tabla de Excel)

openpyxl no sabe crear tablas dinámicas, así que este script las añade escribiendo
directamente el XML del formato .xlsx (un .xlsx es un .zip con archivos XML dentro).
Las tablas dinámicas se marcan como «refreshOnLoad»: Excel las calcula al abrir el libro.

Uso (desde proyecto_vtc/):  python3 scripts/crear_excel.py
"""
import json
import re
import shutil
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).resolve().parents[1]
LIMPIOS = BASE / "datos" / "limpios"
SALIDA = BASE / "analisis_vtc.xlsx"

v = pd.read_csv(LIMPIOS / "viajes_limpios.csv")
j = pd.read_csv(LIMPIOS / "jornadas_limpias.csv")
m = json.loads((LIMPIOS / "metricas_clave.json").read_text())

FUENTE = "Arial"
F = lambda **k: Font(name=FUENTE, **k)
AZUL_TXT = "0000FF"         # dato introducido a mano / valor fijo
AMARILLO = PatternFill("solid", fgColor="FFFF00")
CABECERA = PatternFill("solid", fgColor="DCE6F2")
GRIS = PatternFill("solid", fgColor="F2F2F2")
fino = Side(style="thin", color="BFBFBF")
BORDE = Border(top=fino, bottom=fino, left=fino, right=fino)
EUR = '#,##0.00 "€";-#,##0.00 "€";"-"'
PCT = '0.0%;-0.0%;"-"'

wb = Workbook()
ws_leeme = wb.active
ws_leeme.title = "LEEME"
ws_res = wb.create_sheet("Resumen")
ws_esc = wb.create_sheet("Escenarios")
ws_td = {n: wb.create_sheet(n) for n in ["TD_DiaSemana", "TD_Franjas", "TD_Producto", "TD_Mes_Pago"]}
ws_v = wb.create_sheet("Viajes")
ws_j = wb.create_sheet("Jornadas")

# ---------------------------------------------------------------- Viajes (datos)
COLS = ["id_viaje", "fecha", "hora", "mes", "semana_iso", "dia_semana", "num_dia_semana", "hora_del_dia",
        "franja_horaria", "producto", "importe_eur", "metodo_pago", "sin_ingreso", "importe_6_exactos",
        "es_aeropuerto", "dia_completo"]
datos = v[COLS].copy()
datos["fecha"] = pd.to_datetime(datos.fecha)
ws_v.append(COLS)
for fila in datos.itertuples(index=False):
    ws_v.append([x.to_pydatetime() if isinstance(x, pd.Timestamp) else (bool(x) if isinstance(x, (bool,)) or str(type(x)).endswith("bool_'>") else x)
                 for x in fila])
N = len(datos) + 1  # última fila con datos
# Columna auxiliar con fórmula: 1 en el primer viaje de cada fecha (sirve para contar días distintos con SUMIFS)
AUX = get_column_letter(len(COLS) + 1)
ws_v.cell(1, len(COLS) + 1, "es_primer_viaje_dia")
for r in range(2, N + 1):
    ws_v.cell(r, len(COLS) + 1, f"=IF(COUNTIF($B$2:B{r},B{r})=1,1,0)")
for c in ws_v[1]:
    c.font = F(bold=True)
for r in range(2, N + 1):
    ws_v.cell(r, 2).number_format = "dd/mm/yyyy"
    ws_v.cell(r, 11).number_format = EUR
    for col in range(1, len(COLS) + 2):
        ws_v.cell(r, col).font = F()
tab = Table(displayName="tViajes", ref=f"A1:{AUX}{N}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws_v.add_table(tab)
for i, w in enumerate([9, 11, 7, 9, 10, 11, 8, 8, 13, 18, 11, 11, 11, 11, 11, 11, 12], 1):
    ws_v.column_dimensions[get_column_letter(i)].width = w
ws_v.freeze_panes = "A2"

# Rangos con nombre de columna (para que las fórmulas se lean mejor)
col = {c: get_column_letter(i + 1) for i, c in enumerate(COLS)}
col["es_primer_viaje_dia"] = AUX
R = lambda c: f"Viajes!${col[c]}$2:${col[c]}${N}"

# ---------------------------------------------------------------- Jornadas (datos + fórmulas)
JCOLS = ["fecha", "dia_semana", "inicio", "fin", "conduccion_efectiva_min", "recarga_electrica_min",
         "cuidado_vehiculo_min", "conduccion_no_efectiva_min", "sin_clasificar_min", "tiempo_efectivo_min",
         "tiempo_total_app_min", "jornada_anomala"]
ws_j.append(JCOLS + ["ingresos_eur", "horas_totales", "eur_por_hora_total", "eur_por_hora_efectiva"])
for i, r in enumerate(j[JCOLS].itertuples(index=False), start=2):
    fila = list(r)
    fila[0] = pd.to_datetime(fila[0]).to_pydatetime()
    fila[-1] = bool(fila[-1])
    fila += [f"=SUMIFS({R('importe_eur')},{R('fecha')},A{i})",       # ingresos del día desde la hoja Viajes
             f"=K{i}/60",
             f"=M{i}/N{i}",
             f"=M{i}/(J{i}/60)"]
    ws_j.append(fila)
NJ = len(j) + 1
for c in ws_j[1]:
    c.font = F(bold=True)
for r in range(2, NJ + 1):
    ws_j.cell(r, 1).number_format = "dd/mm/yyyy"
    for c in range(1, 17):
        ws_j.cell(r, c).font = F()
    ws_j.cell(r, 13).number_format = EUR
    ws_j.cell(r, 14).number_format = "0.00"
    ws_j.cell(r, 15).number_format = EUR
    ws_j.cell(r, 16).number_format = EUR
tj = Table(displayName="tJornadas", ref=f"A1:P{NJ}")
tj.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws_j.add_table(tj)
for i in range(1, 17):
    ws_j.column_dimensions[get_column_letter(i)].width = 14
ws_j.freeze_panes = "A2"
ws_j.cell(NJ + 2, 1, "Fuente: app de la flota («Detalle de jornada»). Columnas M–P son fórmulas. "
                     "jornada_anomala = VERDADERO el 02/10 (probable jornada no cerrada).").font = F(italic=True, color="595959")

# ---------------------------------------------------------------- Resumen (fórmulas + gráficos)
ws = ws_res
ws["A1"] = "Resumen del análisis VTC (1/09–7/10/2026)"
ws["A1"].font = F(bold=True, size=14)
ws["A2"] = "Todas las cifras son FÓRMULAS sobre las hojas Viajes y Jornadas. Solo días completos (sin la noche del 31/08)."
ws["A2"].font = F(italic=True, color="595959")

def cab(fila, textos):
    for k, t in enumerate(textos):
        c = ws.cell(fila, 1 + k, t)
        c.font = F(bold=True)
        c.fill = CABECERA
        c.border = BORDE
        c.alignment = Alignment(wrap_text=True, vertical="center")

cab(4, ["Indicador", "Valor", "Cómo se calcula"])
DIAS_DISTINTOS = f"SUMIFS({R('es_primer_viaje_dia')},{R('dia_completo')},TRUE)"
kpis = [
    ("Días completos trabajados", f"={DIAS_DISTINTOS}", "0", "Suma de la columna auxiliar es_primer_viaje_dia (1 por fecha) en días completos"),
    ("Viajes", f"=COUNTIFS({R('dia_completo')},TRUE)", "0", "COUNTIFS"),
    ("Ingresos totales", f"=SUMIFS({R('importe_eur')},{R('dia_completo')},TRUE)", EUR, "SUMIFS de importe_eur"),
    ("Ingreso medio por día", "=B7/B5", EUR, "Ingresos ÷ días"),
    ("Viajes por día", "=B6/B5", "0.0", "Viajes ÷ días"),
    ("Importe medio por viaje con ingreso", f"=AVERAGEIFS({R('importe_eur')},{R('importe_eur')},\">0\",{R('dia_completo')},TRUE)", EUR, "AVERAGEIFS, excluye los viajes de 0 €"),
    ("% de viajes a 0 €", f"=COUNTIFS({R('importe_eur')},0,{R('dia_completo')},TRUE)/B6", PCT, "Probables cancelaciones"),
    ("% de viajes de 6 € exactos", f"=COUNTIFS({R('importe_eur')},6,{R('dia_completo')},TRUE)/B6", PCT, "Probable importe mínimo"),
    ("€ por hora de jornada (bruto)", "=SUMIFS(Jornadas!$M$2:$M$9,Jornadas!$L$2:$L$9,FALSE)/(SUMIFS(Jornadas!$K$2:$K$9,Jornadas!$L$2:$L$9,FALSE)/60)", EUR, "7 jornadas con horas (sin el 02/10): Σ€ ÷ Σhoras"),
    ("€ por hora efectiva (bruto)", "=SUMIFS(Jornadas!$M$2:$M$9,Jornadas!$L$2:$L$9,FALSE)/(SUMIFS(Jornadas!$J$2:$J$9,Jornadas!$L$2:$L$9,FALSE)/60)", EUR, "Igual, con el tiempo efectivo de la flota"),
    ("Minutos/día sin conducción efectiva", "=AVERAGEIFS(Jornadas!$F$2:$F$9,Jornadas!$L$2:$L$9,FALSE)+AVERAGEIFS(Jornadas!$G$2:$G$9,Jornadas!$L$2:$L$9,FALSE)+AVERAGEIFS(Jornadas!$H$2:$H$9,Jornadas!$L$2:$L$9,FALSE)+AVERAGEIFS(Jornadas!$I$2:$I$9,Jornadas!$L$2:$L$9,FALSE)", "0", "Recarga + cuidado + no efectiva + sin clasificar"),
]
for k, (nom, f, fmt, como) in enumerate(kpis, start=5):
    ws.cell(k, 1, nom).font = F()
    c = ws.cell(k, 2, f)
    c.font = F(bold=True)
    c.number_format = fmt
    ws.cell(k, 3, como).font = F(color="595959")
    for cc in range(1, 4):
        ws.cell(k, cc).border = BORDE

# Tabla por día de la semana
F0 = 18
ws.cell(F0 - 1, 1, "Pregunta 2 · Ingresos por día de la semana").font = F(bold=True, size=12)
cab(F0, ["Día", "Días trabajados", "Ingresos", "Ingreso medio por día"])
for k, d in enumerate(["lunes", "martes", "miércoles", "jueves", "viernes"], start=F0 + 1):
    ws.cell(k, 1, d)
    ws.cell(k, 2, f"=SUMIFS({R('es_primer_viaje_dia')},{R('dia_semana')},A{k},{R('dia_completo')},TRUE)").number_format = "0"
    ws.cell(k, 3, f"=SUMIFS({R('importe_eur')},{R('dia_semana')},A{k},{R('dia_completo')},TRUE)").number_format = EUR
    ws.cell(k, 4, f"=C{k}/B{k}").number_format = EUR
    for cc in range(1, 5):
        ws.cell(k, cc).font = F()
        ws.cell(k, cc).border = BORDE

# Tabla por franja horaria
F1 = 27
ws.cell(F1 - 1, 1, "Pregunta 2 · Franjas horarias (hora de inicio del viaje)").font = F(bold=True, size=12)
cab(F1, ["Franja", "Viajes", "Ingresos", "% viajes a 0 €", "€ por hora trabajada*"])
franjas = sorted(v[v.dia_completo].franja_horaria.unique())
for k, fr in enumerate(franjas, start=F1 + 1):
    ws.cell(k, 1, fr)
    ws.cell(k, 2, f"=COUNTIFS({R('franja_horaria')},A{k},{R('dia_completo')},TRUE)")
    ws.cell(k, 3, f"=SUMIFS({R('importe_eur')},{R('franja_horaria')},A{k},{R('dia_completo')},TRUE)").number_format = EUR
    ws.cell(k, 4, f"=COUNTIFS({R('franja_horaria')},A{k},{R('importe_eur')},0,{R('dia_completo')},TRUE)/B{k}").number_format = PCT
    val = m["eur_por_franja"].get(fr[:5])
    c = ws.cell(k, 5, val)
    c.number_format = EUR
    for cc in range(1, 6):
        ws.cell(k, cc).font = F()
        ws.cell(k, cc).border = BORDE
    c.font = F(color=AZUL_TXT)
nota_fr = F1 + len(franjas) + 1
ws.cell(nota_fr, 1, "* Valor fijo (texto azul) calculado en 02_analisis.ipynb: € de cada hora ÷ días en que trabajaste esa hora "
                    "(entre tu primer y tu último viaje). No se puede hacer con una fórmula sencilla.").font = F(italic=True, color="595959")
ws.cell(F1, 5).comment = Comment("Fuente: datos/limpios/metricas_clave.json (02_analisis.ipynb, Pregunta 2).", "Proyecto VTC")

# Tabla por producto
F2 = nota_fr + 3
ws.cell(F2 - 1, 1, "Preguntas 3 y 5 · Tipos de servicio (no hay datos de Uber)").font = F(bold=True, size=12)
cab(F2, ["Producto", "Viajes", "Ingresos", "Importe medio sin los de 0 €"])
prods = v[v.dia_completo].producto.value_counts().index.tolist()
for k, p in enumerate(prods, start=F2 + 1):
    ws.cell(k, 1, p)
    ws.cell(k, 2, f"=COUNTIFS({R('producto')},A{k},{R('dia_completo')},TRUE)")
    ws.cell(k, 3, f"=SUMIFS({R('importe_eur')},{R('producto')},A{k},{R('dia_completo')},TRUE)").number_format = EUR
    ws.cell(k, 4, f"=IFERROR(AVERAGEIFS({R('importe_eur')},{R('producto')},A{k},{R('importe_eur')},\">0\",{R('dia_completo')},TRUE),0)").number_format = EUR
    for cc in range(1, 5):
        ws.cell(k, cc).font = F()
        ws.cell(k, cc).border = BORDE
FIN_PROD = F2 + len(prods)

for letra, ancho in zip("ABCDE", [36, 16, 16, 22, 22]):
    ws.column_dimensions[letra].width = ancho
ws.column_dimensions["C"].width = 44

def grafico(titulo, eje_y, cats, vals, ancla, fmt='#,##0 "€"'):
    ch = BarChart()
    ch.type = "col"
    ch.title = titulo
    ch.y_axis.title = eje_y
    ch.y_axis.numFmt = fmt
    ch.y_axis.majorGridlines = None
    ch.legend = None
    ch.add_data(vals, titles_from_data=False)
    ch.set_categories(cats)
    ch.series[0].graphicalProperties.solidFill = "2A78D6"
    ch.series[0].graphicalProperties.line.noFill = True
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.numFmt = fmt
    ch.width, ch.height = 17, 8
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, ancla)

grafico("Por día de la semana no hay un ganador claro (4–6 días cada uno)", "€ medios por día",
        Reference(ws, min_col=1, min_row=F0 + 1, max_row=F0 + 5), Reference(ws, min_col=4, min_row=F0 + 1, max_row=F0 + 5), "G4")
grafico("Las 8:00 son la hora que más rinde; las 12:00 apenas aportan", "€ por hora trabajada",
        Reference(ws, min_col=1, min_row=F1 + 1, max_row=F1 + len(franjas)), Reference(ws, min_col=5, min_row=F1 + 1, max_row=F1 + len(franjas)), "G21")
grafico("Un viaje de aeropuerto vale ~50 % más que uno normal", "€ medio por viaje (sin los de 0 €)",
        Reference(ws, min_col=1, min_row=F2 + 1, max_row=F2 + 5), Reference(ws, min_col=4, min_row=F2 + 1, max_row=F2 + 5), "G38", fmt='#,##0.00 "€"')

# ---------------------------------------------------------------- Escenarios (neto)
ws = ws_esc
ws["A1"] = "Neto por hora: rellena TUS datos en las celdas amarillas"
ws["A1"].font = F(bold=True, size=14)
ws["A2"] = ("Las capturas no muestran comisión ni costes, así que el neto no se puede calcular sin ti. "
            "No se ha inventado ningún coste: las celdas amarillas están vacías a propósito.")
ws["A2"].font = F(italic=True, color="595959")
filas = [
    ("€ brutos por hora de jornada (de los datos)", "=Resumen!B13", EUR, "Fórmula: enlaza con la hoja Resumen"),
    ("Horas por jornada (media de 7 jornadas)", "=AVERAGEIFS(Jornadas!$N$2:$N$9,Jornadas!$L$2:$L$9,FALSE)", "0.00", "Fórmula: hoja Jornadas"),
    ("Días trabajados por semana (media)", f"=Resumen!B5/{v[v.dia_completo].semana_iso.nunique()}", "0.0", "26 días ÷ 6 semanas con actividad (la 41 está a medias: subestima un poco)"),
    ("Horas por semana", "=B5*B6", "0.0", "Fórmula"),
    ("% de lo facturado que te queda (tras comisión y reparto con la flota)", None, "0%", "RELLENA (ej.: escribe 0,5 para 50 %)"),
    ("Costes fijos por semana (alquiler, cuota, seguro…) en €", None, EUR, "RELLENA (0 si no pagas)"),
    ("Energía/combustible por hora en €", None, EUR, "RELLENA (0 si lo paga la flota)"),
    ("Otros costes por hora en €", None, EUR, "RELLENA (0 si no hay)"),
]
cab_e = ["Concepto", "Valor", "Nota"]
for k, t in enumerate(cab_e):
    c = ws.cell(3, 1 + k, t)
    c.font = F(bold=True)
    c.fill = CABECERA
for k, (nom, f, fmt, nota) in enumerate(filas, start=4):
    ws.cell(k, 1, nom).font = F()
    c = ws.cell(k, 2, f)
    c.number_format = fmt
    c.font = F(bold=True)
    if f is None:
        c.fill = AMARILLO
        c.font = F(bold=True, color=AZUL_TXT)
    ws.cell(k, 3, nota).font = F(color="595959")
ws["A13"] = "NETO por hora (€)"
ws["A13"].font = F(bold=True, size=12)
ws["B13"] = '=IF(COUNT(B8:B11)<4,"Rellena las 4 celdas amarillas",B4*B8-B9/B7-B10-B11)'
ws["B13"].number_format = EUR
ws["B13"].font = F(bold=True, size=12)
ws["A14"] = "NETO por semana (€)"
ws["A14"].font = F(bold=True)
ws["B14"] = '=IF(ISNUMBER(B13),B13*B7,"")'
ws["B14"].number_format = EUR
ws["B14"].font = F(bold=True)

ws["A17"] = "Tabla de escenarios HIPOTÉTICOS (no son datos): neto €/hora según % que te queda (filas) y costes fijos semanales (columnas), sin energía"
ws["A17"].font = F(bold=True)
pcts = [0.4, 0.5, 0.6, 0.7, 0.8, 1.0]
fijos = [0, 100, 200, 300]
ws.cell(18, 1, "% que te queda  \\  costes fijos €/semana").font = F(bold=True)
for k, fj in enumerate(fijos):
    c = ws.cell(18, 2 + k, fj)
    c.font = F(bold=True, color=AZUL_TXT)
    c.number_format = '#,##0 "€"'
    c.fill = GRIS
for r, p in enumerate(pcts, start=19):
    c = ws.cell(r, 1, p)
    c.number_format = "0%"
    c.font = F(bold=True, color=AZUL_TXT)
    c.fill = GRIS
    for k in range(len(fijos)):
        colL = get_column_letter(2 + k)
        cc = ws.cell(r, 2 + k, f"=$B$4*$A{r}-{colL}$18/$B$7")
        cc.number_format = EUR
        cc.font = F()
ws.cell(26, 1, "Los valores de la cabecera de la tabla (texto azul) son supuestos para practicar; cámbialos si quieres.").font = F(italic=True, color="595959")
ws.column_dimensions["A"].width = 62
for L in "BCDE":
    ws.column_dimensions[L].width = 18
ws.column_dimensions["C"].width = 62

# ---------------------------------------------------------------- TD_* (las rellena el XML de abajo)
NOTAS_TD = {
    "TD_DiaSemana": "Tabla dinámica: ingresos y nº de viajes por día de la semana. Filtro: dia_completo = VERDADERO.",
    "TD_Franjas": "Tabla dinámica: viajes, ingresos e importe medio por franja horaria. Filtro: dia_completo = VERDADERO.",
    "TD_Producto": "Tabla dinámica: viajes, ingresos e importe medio por tipo de servicio (todos los viajes).",
    "TD_Mes_Pago": "Tabla dinámica: ingresos por mes y método de pago. Compárala con la app: sep 1.620,42 € app + 122,61 € efectivo; oct 394,87 € + 57,79 €.",
}
for nombre, hoja in ws_td.items():
    hoja["F1"] = NOTAS_TD[nombre]
    hoja["F1"].font = F(italic=True, color="595959")
    hoja["F2"] = "Si sale vacía: pestaña Datos → Actualizar todo. Para practicar: arrastra otros campos en el panel de la derecha."
    hoja["F2"].font = F(italic=True, color="595959")
    hoja.column_dimensions["A"].width = 22
    for L in "BCD":
        hoja.column_dimensions[L].width = 18

# ---------------------------------------------------------------- LEEME
ws = ws_leeme
lineas = [
    ("Análisis VTC · 1/09–7/10/2026", F(bold=True, size=14)),
    ("Datos reales del conductor, anonimizados (sin pasajeros, matrículas ni direcciones).", F(italic=True)),
    ("", F()),
    ("Hojas", F(bold=True, size=12)),
    ("Resumen: indicadores y tablas con fórmulas (SUMIFS, COUNTIFS, AVERAGEIFS) y 3 gráficos.", F()),
    ("Escenarios: calcula tu neto por hora. Rellena las celdas AMARILLAS con tus costes.", F()),
    ("TD_DiaSemana, TD_Franjas, TD_Producto, TD_Mes_Pago: tablas dinámicas reales sobre la hoja Viajes.", F()),
    ("Viajes: 212 viajes (Tabla de Excel «tViajes»). Jornadas: 8 jornadas con horas de la app de la flota.", F()),
    ("", F()),
    ("Leyenda de colores", F(bold=True, size=12)),
    ("Texto negro = fórmula · Texto azul = valor fijo/supuesto · Fondo amarillo = celda para que la rellenes tú.", F()),
    ("", F()),
    ("Cómo practicar con las tablas dinámicas", F(bold=True, size=12)),
    ("1. Haz clic dentro de una tabla dinámica → aparece el panel «Campos de tabla dinámica».", F()),
    ("2. Arrastra campos a Filas, Columnas, Valores o Filtros. Prueba: Filas = semana_iso, Valores = Suma de importe_eur.", F()),
    ("3. Clic derecho en un valor → «Configuración de campo de valor» para cambiar Suma por Promedio o Cuenta.", F()),
    ("4. Si cambias datos en la hoja Viajes: Datos → Actualizar todo.", F()),
    ("", F()),
    ("Limitaciones principales: solo 8 jornadas con horas; sin zonas, sin Uber, sin costes. Ver LIMITACIONES.md.", F(italic=True)),
]
for k, (t, fnt) in enumerate(lineas, start=1):
    ws.cell(k, 1, t).font = fnt
ws.column_dimensions["A"].width = 120

wb.calculation.fullCalcOnLoad = True   # Excel recalcula todas las fórmulas al abrir
wb.save(SALIDA)

# =====================================================================
#  TABLAS DINÁMICAS: se añaden editando el XML dentro del .xlsx
# =====================================================================
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NSR = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_PT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable"
REL_PC = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition"
CT_PT = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"
CT_PC = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"

idx = {c: i for i, c in enumerate(COLS)}
ORDEN_DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
items = {   # valores distintos de los campos que irán en filas/columnas/filtros (en el orden en que se mostrarán)
    "dia_semana": [d for d in ORDEN_DIAS if d in set(datos.dia_semana)],
    "franja_horaria": sorted(datos.franja_horaria.unique()),
    "producto": datos.producto.value_counts().index.tolist(),
    "mes": sorted(datos.mes.unique()),
    "metodo_pago": sorted(datos.metodo_pago.unique()),
}
BOOL = {"sin_ingreso", "importe_6_exactos", "es_aeropuerto", "dia_completo"}


def cache_field(c):
    s = datos[c]
    if c in items:
        inner = "".join(f'<s v="{escape(x)}"/>' for x in items[c])
        return f'<cacheField name="{c}" numFmtId="0"><sharedItems count="{len(items[c])}">{inner}</sharedItems></cacheField>'
    if c in BOOL:
        return f'<cacheField name="{c}" numFmtId="0"><sharedItems count="2"><b v="0"/><b v="1"/></sharedItems></cacheField>'
    if c == "fecha":
        return (f'<cacheField name="fecha" numFmtId="14"><sharedItems containsSemiMixedTypes="0" containsNonDate="0" '
                f'containsDate="1" containsString="0" minDate="{s.min():%Y-%m-%dT00:00:00}" maxDate="{s.max():%Y-%m-%dT00:00:00}"/></cacheField>')
    if pd.api.types.is_numeric_dtype(s):
        entero = ' containsInteger="1"' if pd.api.types.is_integer_dtype(s) else ""
        return (f'<cacheField name="{c}" numFmtId="0"><sharedItems containsSemiMixedTypes="0" containsString="0" '
                f'containsNumber="1"{entero} minValue="{s.min()}" maxValue="{s.max()}"/></cacheField>')
    return f'<cacheField name="{c}" numFmtId="0"><sharedItems/></cacheField>'


cache_xml = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
             f'<pivotCacheDefinition xmlns="{NS}" xmlns:r="{NSR}" saveData="0" refreshOnLoad="1" createdVersion="6" '
             f'refreshedVersion="6" minRefreshableVersion="3" recordCount="{len(datos)}">'
             f'<cacheSource type="worksheet"><worksheetSource ref="A1:{get_column_letter(len(COLS))}{N}" sheet="Viajes"/></cacheSource>'
             f'<cacheFields count="{len(COLS)}">' + "".join(cache_field(c) for c in COLS) + '</cacheFields></pivotCacheDefinition>')

# Definición de cada tabla dinámica: filas, columnas, filtro y valores
TDS = [
    dict(hoja="TD_DiaSemana", nombre="TD_DiaSemana", filas="dia_semana", columnas=None, filtro="dia_completo",
         valores=[("Suma de importe_eur", "importe_eur", "sum", 4), ("Nº de viajes", "id_viaje", "count", 0)]),
    dict(hoja="TD_Franjas", nombre="TD_Franjas", filas="franja_horaria", columnas=None, filtro="dia_completo",
         valores=[("Nº de viajes", "id_viaje", "count", 0), ("Suma de importe_eur", "importe_eur", "sum", 4),
                  ("Promedio de importe_eur", "importe_eur", "average", 4)]),
    dict(hoja="TD_Producto", nombre="TD_Producto", filas="producto", columnas=None, filtro=None,
         valores=[("Nº de viajes", "id_viaje", "count", 0), ("Suma de importe_eur", "importe_eur", "sum", 4),
                  ("Promedio de importe_eur", "importe_eur", "average", 4)]),
    dict(hoja="TD_Mes_Pago", nombre="TD_Mes_Pago", filas="mes", columnas="metodo_pago", filtro=None,
         valores=[("Suma de importe_eur", "importe_eur", "sum", 4)]),
]


def pivot_xml(td):
    fila_ini = 3 if td["filtro"] else 1                    # con filtro, el filtro va en A1 y la tabla empieza en A3
    n_items = len(items[td["filas"]])
    n_val = len(td["valores"])
    if td["columnas"]:
        n_col_items = len(items[td["columnas"]])
        alto = 2 + n_items + 1
        ancho = 1 + n_col_items + 1
        first = 'firstHeaderRow="1" firstDataRow="2" firstDataCol="1"'
    else:
        alto = 1 + n_items + 1
        ancho = 1 + n_val
        first = f'firstHeaderRow="{0 if n_val > 1 else 1}" firstDataRow="1" firstDataCol="1"'
    ref = f"A{fila_ini}:{get_column_letter(ancho)}{fila_ini + alto - 1}"
    page = ' rowPageCount="1" colPageCount="1"' if td["filtro"] else ""

    pfs = []
    datos_campos = {v[1] for v in td["valores"]}
    for c in COLS:
        attrs = ' dataField="1"' if c in datos_campos else ""
        if c == td["filas"] or c == td["columnas"]:
            eje = "axisRow" if c == td["filas"] else "axisCol"
            its = "".join(f'<item x="{k}"/>' for k in range(len(items[c]))) + '<item t="default"/>'
            pfs.append(f'<pivotField axis="{eje}"{attrs} showAll="0"><items count="{len(items[c]) + 1}">{its}</items></pivotField>')
        elif c == td["filtro"]:
            pfs.append(f'<pivotField axis="axisPage"{attrs} showAll="0"><items count="3"><item x="0"/><item x="1"/><item t="default"/></items></pivotField>')
        else:
            pfs.append(f'<pivotField{attrs} showAll="0"/>')
    x = '<pivotFields count="%d">%s</pivotFields>' % (len(COLS), "".join(pfs))
    x += f'<rowFields count="1"><field x="{idx[td["filas"]]}"/></rowFields>'
    x += f'<rowItems count="{n_items + 1}">' + "".join(
        "<i><x/></i>" if k == 0 else f'<i><x v="{k}"/></i>' for k in range(n_items)) + '<i t="grand"><x/></i></rowItems>'
    if td["columnas"]:
        x += f'<colFields count="1"><field x="{idx[td["columnas"]]}"/></colFields>'
        x += f'<colItems count="{n_col_items + 1}">' + "".join(
            "<i><x/></i>" if k == 0 else f'<i><x v="{k}"/></i>' for k in range(n_col_items)) + '<i t="grand"><x/></i></colItems>'
    elif n_val > 1:
        x += '<colFields count="1"><field x="-2"/></colFields>'
        x += f'<colItems count="{n_val}">' + "".join(
            "<i><x/></i>" if k == 0 else f'<i i="{k}"><x v="{k}"/></i>' for k in range(n_val)) + '</colItems>'
    else:
        x += '<colItems count="1"><i/></colItems>'
    if td["filtro"]:
        # item="1" -> el segundo elemento de la lista (VERDADERO): solo días completos
        x += f'<pageFields count="1"><pageField fld="{idx[td["filtro"]]}" item="1" hier="-1"/></pageFields>'
    x += f'<dataFields count="{n_val}">' + "".join(
        f'<dataField name="{escape(n)}" fld="{idx[c]}" subtotal="{s}" baseField="0" baseItem="0" numFmtId="{nf}"/>'
        for n, c, s, nf in td["valores"]) + '</dataFields>'
    x += '<pivotTableStyleInfo name="PivotStyleMedium2" showRowHeaders="1" showColHeaders="1" showRowStripes="0" showColStripes="0" showLastColumn="1"/>'
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<pivotTableDefinition xmlns="{NS}" name="{td["nombre"]}" cacheId="1" applyNumberFormats="0" applyBorderFormats="0" '
            f'applyFontFormats="0" applyPatternFormats="0" applyAlignmentFormats="0" applyWidthHeightFormats="1" '
            f'dataCaption="Valores" updatedVersion="6" minRefreshableVersion="3" useAutoFormatting="1" itemPrintTitles="1" '
            f'createdVersion="6" indent="0" outline="1" outlineData="1" multipleFieldFilters="0" '
            f'rowHeaderCaption="{escape(td["filas"])}">'
            f'<location ref="{ref}" {first}{page}/>' + x + '</pivotTableDefinition>')


tmp = SALIDA.with_suffix(".tmp.xlsx")
with zipfile.ZipFile(SALIDA) as zin:
    archivos = {n: zin.read(n) for n in zin.namelist()}

wbxml = archivos["xl/workbook.xml"].decode()
wbrels = archivos["xl/_rels/workbook.xml.rels"].decode()
ctypes = archivos["[Content_Types].xml"].decode()

# hoja -> archivo sheetN.xml (según workbook.xml y sus relaciones)
rid_de_hoja = dict(re.findall(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbxml))
destino = dict(re.findall(r'Id="([^"]+)"[^>]*Target="/?(?:xl/)?([^"]+)"', wbrels)) | \
          {a: b for b, a in re.findall(r'Target="/?(?:xl/)?([^"]+)"[^>]*Id="([^"]+)"', wbrels)}

archivos["xl/pivotCache/pivotCacheDefinition1.xml"] = cache_xml.encode()
wbrels = wbrels.replace("</Relationships>",
                        f'<Relationship Id="rIdPC1" Type="{REL_PC}" Target="pivotCache/pivotCacheDefinition1.xml"/></Relationships>')
pc = '<pivotCaches><pivotCache cacheId="1" r:id="rIdPC1"/></pivotCaches>'
if "<calcPr" in wbxml:
    wbxml = re.sub(r"(<calcPr[^>]*/>)", r"\1" + pc, wbxml, count=1)
else:
    wbxml = wbxml.replace("</sheets>", "</sheets>" + pc)
ctypes = ctypes.replace("</Types>", f'<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="{CT_PC}"/></Types>')

for k, td in enumerate(TDS, start=1):
    hoja_xml = destino[rid_de_hoja[td["hoja"]]]                  # p. ej. worksheets/sheet4.xml
    rels_path = "xl/" + hoja_xml.replace("worksheets/", "worksheets/_rels/") + ".rels"
    archivos[f"xl/pivotTables/pivotTable{k}.xml"] = pivot_xml(td).encode()
    archivos[f"xl/pivotTables/_rels/pivotTable{k}.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="{REL_PC}" Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>').encode()
    rel = f'<Relationship Id="rIdPT{k}" Type="{REL_PT}" Target="../pivotTables/pivotTable{k}.xml"/>'
    if rels_path in archivos:
        archivos[rels_path] = archivos[rels_path].decode().replace("</Relationships>", rel + "</Relationships>").encode()
    else:
        archivos[rels_path] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                               f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rel}</Relationships>').encode()
    ctypes = ctypes.replace("</Types>", f'<Override PartName="/xl/pivotTables/pivotTable{k}.xml" ContentType="{CT_PT}"/></Types>')

archivos["xl/workbook.xml"] = wbxml.encode()
archivos["xl/_rels/workbook.xml.rels"] = wbrels.encode()
archivos["[Content_Types].xml"] = ctypes.encode()
with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for n in ["[Content_Types].xml"] + [n for n in archivos if n != "[Content_Types].xml"]:
        zout.writestr(n, archivos[n])
shutil.move(tmp, SALIDA)
print(f"Creado {SALIDA.name}: {len(TDS)} tablas dinámicas, hojas: {wb.sheetnames}")
